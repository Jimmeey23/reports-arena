# Companion workbook for the three September-2026 Bandra reports (Sales / Class / Instructor)
import pandas as pd, numpy as np
import d1_data as D
import d2_extra as X
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = '/home/user/reports-arena/Bandra_September_2026_Reports_Data.xlsx'
GOLD = '8C6A2F'; DARK = '2E2A26'
hfill = PatternFill('solid', fgColor=GOLD)
hfont = Font(bold=True, color='FFFFFF', size=10)
tfont = Font(bold=True, size=13, color=DARK)
thin = Side(style='thin', color='D9D2C6')
bd = Border(bottom=thin)

wb = Workbook(); wb.remove(wb.active)

def sheet(name, title, note, frames):
    ws = wb.create_sheet(name)
    r = 1
    ws.cell(r, 1, title).font = tfont; r += 1
    ws.cell(r, 1, note).font = Font(italic=True, size=9, color='6B6459'); r += 2
    for head, df, extra in frames:
        if head:
            ws.cell(r, 1, head).font = Font(bold=True, size=11, color=DARK); r += 1
        d = df.copy()
        if d.index.name or not isinstance(d.index, pd.RangeIndex):
            d = d.reset_index()
        for j, c in enumerate(d.columns, start=1):
            cc = ws.cell(r, j, str(c)); cc.fill = hfill; cc.font = hfont
            cc.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.freeze_panes = ws.cell(r + 1, 1)
        r += 1
        for _, row in d.iterrows():
            for j, c in enumerate(d.columns, start=1):
                v = row[c]
                try:
                    if isinstance(v, pd.Period): v = str(v)
                    elif isinstance(v, pd.Timestamp): v = str(v.date())
                    elif isinstance(v, np.generic): v = v.item()
                    elif hasattr(v, 'item') and not isinstance(v, (str, bytes)):
                        v = v.item()
                except Exception:
                    v = str(v)
                if isinstance(v, float): v = round(v, 2)
                ws.cell(r, j, v).border = bd
            r += 1
        r += 1
        for j, c in enumerate(d.columns, start=1):
            w = max([len(str(c))] + [len(str(x)) for x in d[c].tolist()[:60]]) + 3
            cur = ws.column_dimensions[get_column_letter(j)].width or 0
            ws.column_dimensions[get_column_letter(j)].width = max(cur, min(w, 46))
    return ws

R = lambda x: float(x)
def money(df, cols):
    d = df.copy()
    for c in cols:
        if c in d: d[c] = d[c].apply(lambda v: R(v) if isinstance(v, (int, float, np.number)) else v)
    return d

# ---------- 1. SALES ----------
d = D.DN
kpi = pd.DataFrame([
    dict(measure='Cash received', y2024=d['s_cash'][2024], y2025=d['s_cash'][2025], y2026=d['s_cash'][2026],
         yoy_pct=round((d['s_cash'][2026]/d['s_cash'][2025]-1)*100, 2)),
    dict(measure='Transactions', y2024=np.nan, y2025=d['s_txn'][2025], y2026=d['s_txn'][2026],
         yoy_pct=round((d['s_txn'][2026]/d['s_txn'][2025]-1)*100, 2)),
    dict(measure='Paying members', y2024=np.nan, y2025=d['s_mem'][2025], y2026=d['s_mem'][2026],
         yoy_pct=round((d['s_mem'][2026]/d['s_mem'][2025]-1)*100, 2)),
    dict(measure='Average sale', y2024=np.nan, y2025=round(d['s_avg'][2025], 0), y2026=round(d['s_avg'][2026], 0),
         yoy_pct=round((d['s_avg'][2026]/d['s_avg'][2025]-1)*100, 2)),
    dict(measure='Median sale', y2024=np.nan, y2025=d['s_med'][2025], y2026=d['s_med'][2026], yoy_pct=np.nan),
    dict(measure='Discount given', y2024=np.nan, y2025=np.nan, y2026=float(d['s_disc'][2026]), yoy_pct=np.nan),
])
ws = sheet('1 Sales', 'Sales performance — Supreme HQ, Bandra — September',
           'Sourced from Sales - Raw, Day End Part 3 Sessions and the member file. Figures in INR.',
           [('Headline — September 2024 vs 2025 vs 2026', kpi, None),
            ('Category revenue by September', D.CATS, None),
            ('Category revenue, monthly Jan-2025 to Sep-2026', D.CATTR, None),
            ('Product detail, September 2026 vs 2025', D.PRODS, None),
            ('Largest movers in revenue (>= ₹10,000 either year)', D.MOVERS, None),
            ('Sale-size bands, September 2026', D.BANDS, None),
            ('Price realised per class, by product', D.PPC, None),
            ('Sold-by team member, September 2026 vs 2025', D.SOLD, None),
            ('Payment channel, September 2026', X.PM, None),
            ('Discount codes used, September 2026',
             X.sale26.groupby('Discount Code').agg(sales=('Sale ID','nunique'), revenue=('val','sum'),
                 discount=('disc','sum')).sort_values('revenue', ascending=False), None),
            ('Buyer concentration — top 20 by spend, September 2026',
             X.spend.nlargest(20).rename('cash').to_frame(), None),
            ('Members, attendances and revenue per member, by month', X.ACT, None),
            ('The two September core cohorts, like for like', pd.DataFrame([
                dict(cohort='Sept-2024 → Sept-2025 (present in both)', members=len(X.core_old),
                     earlier_visits=round(X.core_old_prev.sum(),0), later_visits=round(X.core_old.sum(),0),
                     earlier_each=round(X.core_old_prev.mean(),2), later_each=round(X.core_old.mean(),2)),
                dict(cohort='Sept-2025 → Sept-2026 (present in both)', members=len(X.core_ix),
                     earlier_visits=round(X.core_then.sum(),0), later_visits=round(X.core_now.sum(),0),
                     earlier_each=round(X.core_then.mean(),2), later_each=round(X.core_now.mean(),2))]), None)])

# ---------- 2. CLASS ----------
sheet('2 Classes', 'Class performance — Supreme HQ, Bandra — September',
      'Sourced from Day End Part 3 - Sessions. Capacity is per session; seats offered = capacity x sessions.',
      [('Format scorecard, September 2026', D.FMT, None),
       ('Format year-on-year comparison', D.FMT_YOY, None),
       ('Day-of-week pattern', D.DOW, None),
       ('Hour-of-day pattern', D.HOUR, None),
       ('Time-band pattern', D.BAND, None),
       ('Membership vs package vs single-class mix by format', D.PAYMIX, None),
       ('Thin sessions — empty and 1–2 attendees, by format', D.LEAN_BY, None),
       ('Fullest sessions of the month', D.FULLEST, None),
       ('Week-by-week fade', pd.concat([v.assign(year=y) for y, v in D.WK.items()]).reset_index()
          .rename(columns={'att': 'attended', 'fill': 'fill_pct', 'rev': 'revenue'}), None),
       ('Booking book — booked, attended, cancelled, comp, 2025 vs 2026', pd.DataFrame([
            dict(measure='Seats on the board', sept2025=round(X.x25['cap'].sum(),0), sept2026=round(X.x26['cap'].sum(),0)),
            dict(measure='Booked', sept2025=round(X.x25['Booked'].sum(),0), sept2026=round(X.x26['Booked'].sum(),0)),
            dict(measure='Attended', sept2025=round(X.x25['CheckedIn'].sum(),0), sept2026=round(X.x26['CheckedIn'].sum(),0)),
            dict(measure='Late cancelled', sept2025=round(X.x25['LateCancelled'].sum(),0), sept2026=round(X.x26['LateCancelled'].sum(),0)),
            dict(measure='Complimentary seats', sept2025=round(X.x25['Complimentary'].sum(),0), sept2026=round(X.COMP26,0)),
            dict(measure='Booked as % of seats on the board', sept2025=round(X.BOOK_BOARD_25,1), sept2026=round(X.BOOK_BOARD_26,1)),
            dict(measure='Of seats booked, % attended', sept2025=round(X.ATT_OF_BOOK_25,1), sept2026=round(X.ATT_OF_BOOK_26,1)),
            dict(measure='Class revenue', sept2025=round(X.x25['Revenue'].sum(),0), sept2026=round(X.x26['Revenue'].sum(),0)),
            dict(measure='Revenue per booked seat', sept2025=round(X.x25['Revenue'].sum()/X.x25['Booked'].sum(),0),
                 sept2026=round(X.x26['Revenue'].sum()/X.x26['Booked'].sum(),0))]), None),
       ('Capacity band — how full each room size ran',
        pd.DataFrame([dict(band=b,
            sessions_2026=int(X.CB26.loc[b,'n']) if b in X.CB26.index else 0,
            attended_2026=int(X.CB26.loc[b,'att']) if b in X.CB26.index else 0,
            fill_2026=round(X.CB_FILL_26.get(b, float('nan')),1),
            sessions_2025=int(X.CB25.loc[b,'n']) if b in X.CB25.index else 0,
            fill_2025=round(X.CB_FILL_25.get(b, float('nan')),1),
            revenue_2026=round(X.CB_REV26.get(b, 0),0)) for b in ['10–11','12–13','14+']]), None),
       ('Complimentary seats by format, September 2026',
        (X.x26.groupby('Class').agg(attended=('CheckedIn','sum'), comp=('Complimentary','sum'))
          .assign(comp_pct=lambda d: (d['comp']/d['attended'].replace(0,float('nan'))*100).round(1))
          .sort_values('comp', ascending=False)), None),
       ('First class circumstances → lifetime depth (members first seen from Jan-2025)',
        pd.concat([X.Q_DAY.assign(by='Day of first class').reset_index().rename(columns={'Day of Week':'value'}),
                   X.Q_BAND.assign(by='Hour of first class').reset_index().rename(columns={'band':'value'}),
                   X.Q_FMT.assign(by='Format of first class').reset_index().rename(columns={'Cleaned Class':'value'}),
                   X.Q_ENT.assign(by='Entry type').reset_index().rename(columns={'Is New':'value'})])
        [['by','value','n','r3','r6','av']].rename(columns={'n':'members','r3':'reached_3rd_pct','r6':'reached_6th_pct','av':'avg_visits'}), None),
       ('Lead source performance, Sept-2025 to Aug-2026 intake', X.LS, None),
       ('Conversion speed buckets, Sept-2025 to Aug-2026 intake', X.SPEED.rename('members').to_frame(), None)])

# ---------- 3. INSTRUCTOR ----------
wb3frames = [
    ('Trainer scorecard, September 2026 vs 2025', D.TCH, None),
    ('Trainer x format', D.TF, None),
    ('Trainer x slot', D.TSLOT, None),
    ('Clients attributed to each trainer', D.CLI, None),
    ('September-2026 client cohort by trainer', D.CLI_SEP, None),
    ('Active recurring slots (from 1 Oct 2026)', D.ACT_SLOTS, None),
    ('Last available payroll — August 2026', D.PAY[['Teacher Name', 'Total Sessions', 'Total Paid']], None),
    ('Payroll history by trainer', D.PAY_HIST, None),
    ('Payroll trend by month (studio level)', D.PAY_TREND, None),
    ('Same slot, different trainer — the controlled test',
     X.SLOTS.rename(columns={'slot':'Slot','fmt':'Format','k':'Trainers','n':'Sessions','lo':'Worst fill %',
        'hi':'Best fill %','best':'Best trainer','worst':'Worst trainer','spread':'Spread (pts)'}), None),
    ('What the studio pays per session, by discipline, August 2026', pd.DataFrame([
        dict(discipline='Cycle', sessions=round(X.PA['Cycle Sessions'].sum(),0),
             paid=round(X.PA['Cycle Paid'].sum(),0), per_session=round(X.CYCLE_RATE,0)),
        dict(discipline='Barre', sessions=round(X.PA['Barre Sessions'].sum(),0),
             paid=round(X.PA['Barre Paid'].sum(),0), per_session=round(X.BARRE_RATE,0)),
        dict(discipline='All', sessions=round(X.PA['Total Sessions'].sum(),0),
             paid=round(X.PA['Total Paid'].sum(),0), per_session=round(X.PAY_SESS,0))]), None),
    ('Trainer-level per-session pay, August 2026',
     X.PA.assign(per_session=lambda d: (d['Total Paid']/d['Total Sessions'].replace(0,float('nan'))).round(0))
       [['Teacher Name','Total Sessions','Total Empty Sessions','Total Customers','Total Paid','per_session']]
       .sort_values('Total Paid', ascending=False), None),
    ('First-class trainer → does the visitor reach a third visit', X.Q_TR, None),
]
sheet('3 Instructors', 'Instructor performance — Supreme HQ, Bandra — September',
      'Sourced from Day End Part 3 - Sessions, Teacher Recurring, member file and Part 1 - Payroll.'
      ' Payroll ends August 2026; it is included as context only.', wb3frames)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=1, max_row=3):
        for c in row: c.alignment = Alignment(vertical='center', wrap_text=False)
wb.save(OUT)
print('workbook written', OUT)
for ws in wb.worksheets:
    print(' -', ws.title, ws.max_row, 'rows x', ws.max_column, 'cols')
